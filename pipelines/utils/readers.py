# Databricks notebook source
# readers - reusable file discovery and reading helpers.

import os
import re
import subprocess
import sys

from pyspark.sql.types import StringType, StructField, StructType


def _normalize_ext(ext: str) -> str:
    ext = (ext or "").strip().lower()
    return ext if ext.startswith(".") else f".{ext}"


def _dbfs_dir(path: str) -> str:
    if path.startswith("/Volumes/"):
        return path.replace("/Volumes/", "dbfs:/Volumes/")
    return path


def list_files_by_extension(dir_path: str, extensions, patterns=None):
    exts = {_normalize_ext(ext) for ext in extensions}
    pattern_list = [p.lower() for p in (patterns or [])]
    matches = []

    try:
        files = dbutils.fs.ls(_dbfs_dir(dir_path))
    except Exception:
        files = []

    for file_info in files:
        name = file_info.name or os.path.basename(file_info.path)
        lower_name = name.lower()

        if not any(lower_name.endswith(ext) for ext in exts):
            continue

        if pattern_list and not any(pattern in lower_name for pattern in pattern_list):
            continue

        matches.append(file_info.path.replace("dbfs:/Volumes/", "/Volumes/"))

    return sorted(matches)


def list_files_by_extension_recursive(dir_path: str, extensions, patterns=None):
    exts = {_normalize_ext(ext) for ext in extensions}
    pattern_list = [p.lower() for p in (patterns or [])]
    matches = []

    def _walk(path: str):
        try:
            entries = dbutils.fs.ls(_dbfs_dir(path))
        except Exception:
            return

        for entry in entries:
            name = entry.name or os.path.basename(entry.path)
            lower_name = name.lower()

            if lower_name.endswith("/"):
                _walk(entry.path.replace("dbfs:/Volumes/", "/Volumes/"))
                continue

            if not any(lower_name.endswith(ext) for ext in exts):
                continue

            if pattern_list and not any(pattern in lower_name for pattern in pattern_list):
                continue

            matches.append(entry.path.replace("dbfs:/Volumes/", "/Volumes/"))

    _walk(dir_path)
    return sorted(matches)


def pick_input_file(dir_path: str, extensions, patterns=None, required=True):
    matches = list_files_by_extension(dir_path, extensions, patterns=patterns)
    if matches:
        selected = matches[-1]
        print(f"[OK] Arquivo selecionado em {dir_path}: {selected}")
        return selected

    if required:
        ext_list = ", ".join(sorted({_normalize_ext(ext) for ext in extensions}))
        pattern_desc = ", ".join(patterns or [])
        raise FileNotFoundError(
            f"Nenhum arquivo encontrado em {dir_path} para extensoes [{ext_list}]"
            + (f" com filtros [{pattern_desc}]" if pattern_desc else "")
        )

    return None


def read_csv_auto(path: str):
    """Read a CSV file trying common delimiters until one produces a plausible schema."""
    tried = []
    best_df = None
    best_sep = None
    best_width = 0

    for sep in [",", ";", "|", "\t"]:
        df = (
            spark.read
            .option("header", "true")
            .option("sep", sep)
            .option("encoding", "UTF-8")
            .option("quote", '"')
            .option("escape", '"')
            .csv(path)
        )
        cols = [c.strip() for c in df.columns]
        width = len(cols)
        tried.append((sep, cols))

        if width > best_width:
            best_df = df.toDF(*cols)
            best_sep = sep
            best_width = width

        if width > 1:
            df = df.toDF(*cols)
            print(f"[OK] {path} read with separator '{sep}'. Columns: {df.columns}")
            return df

    raise Exception(
        f"Could not detect delimiter for {path}. "
        f"Best guess separator='{best_sep}' columns={best_df.columns if best_df is not None else []} "
        f"All tries={tried}"
    )


def read_xlsx_auto(path: str, sheet_name=None):
    try:
        from openpyxl import load_workbook
    except ImportError:
        subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl"])
        from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return spark.createDataFrame([], schema="dummy string").drop("dummy")

    header = [str(v).strip() if v is not None else "" for v in rows[0]]
    data_rows = rows[1:]
    spark_rows = [
        {header[i]: (None if i >= len(row) or row[i] is None else str(row[i]).strip()) for i in range(len(header))}
        for row in data_rows
        if any(v not in (None, "") for v in row)
    ]

    if not spark_rows:
        return spark.createDataFrame([], schema="dummy string").drop("dummy")

    schema = StructType([StructField(col_name, StringType(), True) for col_name in header])
    return spark.createDataFrame(spark_rows, schema=schema)
